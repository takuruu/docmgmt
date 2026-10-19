"""Extraktion von Text aus verschiedenen Dokumentformaten."""
from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Optional

TEXT_SUFFIXES = {".txt", ".md", ".rst", ".csv", ".json", ".xml", ".html", ".log", ".yml", ".yaml"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff"}
MAX_TEXT_CHARS = 12000


def extract_text(path: Path, mistral_client=None) -> Optional[str]:
    """Extrahiert Text aus einer Datei.

    Für Bilder wird optional ein Mistral-Client für OCR genutzt.
    Gibt None zurück, wenn kein Text extrahiert werden kann.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in IMAGE_SUFFIXES:
        return _extract_text_from_image(path, mistral_client)
    if suffix == ".pdf":
        return _extract_text_from_pdf(path)
    if suffix == ".docx":
        return _extract_text_from_docx(path)
    if suffix in TEXT_SUFFIXES:
        return _read_text_file(path)
    return None


def _read_text_file(path: Path) -> Optional[str]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return None
    return text[:MAX_TEXT_CHARS]


def _extract_text_from_pdf(path: Path) -> Optional[str]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return None
    try:
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception:
        return None
    text = "\n".join(pages).strip()
    return text[:MAX_TEXT_CHARS] if text else None


def _extract_text_from_docx(path: Path) -> Optional[str]:
    try:
        from docx import Document
    except ImportError:
        return None
    try:
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
    except Exception:
        return None
    text = "\n".join(parts).strip()
    return text[:MAX_TEXT_CHARS] if text else None


def _extract_text_from_image(path: Path, mistral_client) -> Optional[str]:
    if mistral_client is None:
        return None
    import base64

    try:
        data = base64.b64encode(path.read_bytes()).decode("utf-8")
        mime = _image_mime(path)
        response = mistral_client.ocr.process(
            model="mistral-ocr-latest",
            pages=[{"type": "image_url", "image_url": f"data:{mime};base64,{data}"}],
        )
        pages = response.pages
    except Exception:
        return None
    text = "\n".join(getattr(page, "text", "") or "" for page in pages).strip()
    return text[:MAX_TEXT_CHARS] if text else None


def _image_mime(path: Path) -> str:
    suffix = path.suffix.lower()
    return {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".webp": "image/webp",
        ".gif": "image/gif",
        ".bmp": "image/bmp",
        ".tiff": "image/tiff",
    }.get(suffix, "application/octet-stream")
